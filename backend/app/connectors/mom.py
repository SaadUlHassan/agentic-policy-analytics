"""MOM Statistics connector: release-dated Excel workbooks (CSV as fallback) from the time-series page.

MOM renames its files at every release (e.g. mrsd_3_qtly_emp_chg_by_ind_21092026.xlsx), so each
fetch discovers the newest file for the dataset's stable prefix on the index page. If the page is
unreachable or no longer links the dataset, the last-known file is used and that is disclosed.

The workbook is preferred: it publishes sector totals with their SSIC codes, its legend and its
unit. The CSV lists only the lowest-level industries (no Information & Communications total), so it
is used only when the workbook is unavailable or no longer has the expected structure.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from typing import Any
from urllib.parse import urljoin

import openpyxl

from app.connectors.base import assemble
from app.connectors.catalog import DatasetEntry, PlatformInfo
from app.connectors.http import Fetcher, RawPayload, SourceSchemaError, SourceUnavailable
from app.connectors.ratelimit import RateLimit
from app.connectors.tabular import read_csv, require_columns
from app.data.models import NormalizedDataset, QualityIssue, SeriesData, Severity
from app.data.series import build_series
from app.data.taxonomy import Taxonomy

INDEX_URL = "https://stats.mom.gov.sg/Pages/EmploymentTimeSeries.aspx"
# MOM publishes no limit; one request a second keeps us a polite client.
POLITE = RateLimit("mom", 1, {"none": 1})

_HREF = re.compile(r'href="([^"]+\.(?:xlsx|csv))"', re.I)
_RELEASE = re.compile(r"_(\d{2})(\d{2})(\d{4})\.(?:xlsx|csv)$", re.I)
_LEGEND = re.compile(r"^'(?P<marker>[^']+)'\s*:\s*(?P<meaning>.+)$")
# Any "In <unit>" note: a unit we do not recognise must block use, not pass as unstated.
_UNIT = re.compile(r"^in\s+(?P<unit>[a-z' ]{1,20})$", re.I)
_UNIT_SCALES = {"thousand": 1000, "hundred": 100, "number": 1, "person": 1}
_HEADER_MARKS = re.compile(r"[\^*#]+$")
BLOCK_HEADER = "SSIC 2020"


def release_date(url: str) -> str:
    """MOM file names end in a DDMMYYYY release date; returned as ISO YYYY-MM-DD ("" if absent)."""
    m = _RELEASE.search(url)
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else ""


@dataclass(frozen=True)
class _RowId:
    """A source row is its stated SSIC code plus its label: two rows sharing a label (e.g. a child
    relabelled with its parent's name) stay separate and are caught by the code check."""

    code: str | None
    label: str

    @property
    def ref(self) -> str:
        return f"{self.code} {self.label}" if self.code else self.label


@dataclass
class _Rows:
    """Source rows and their (period, raw value) points."""

    points: dict[_RowId, list[tuple[str, str]]] = field(default_factory=lambda: defaultdict(list))
    issues: list[QualityIssue] = field(default_factory=list)
    footnotes: str | None = None
    # Safety-relevant statements the sheet makes about itself; part of the artifact's identity.
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class _Release:
    url: str
    fallback_reason: str | None = None  # set when this is the last-known file, not a discovered one

    def disclosure(self, ext: str) -> QualityIssue | None:
        if self.fallback_reason is None:
            return None
        return QualityIssue(
            code="DISCOVERY_FALLBACK",
            severity=Severity.WARNING,
            message=f"{self.fallback_reason}; a newer release may exist. Using the {ext} released "
            f"{release_date(self.url) or 'on an unknown date'}",
        )


async def _discover(entry: DatasetEntry, fetcher: Fetcher) -> dict[str, _Release]:
    stale: str | None = None  # why links came from a stored copy of the page, if they did
    try:
        page = await fetcher.get(INDEX_URL, rate=POLITE, accept="text/html")
        links = [_absolute(h) for h in _HREF.findall(page.body.decode("utf-8", "replace"))]
        reason = "the index page no longer links this dataset"
        if page.fallback_reason:
            stale = (
                f"index page unavailable ({page.fallback_reason}); release links taken from its "
                f"snapshot of {page.retrieved_at.date().isoformat()}"
            )
    except SourceUnavailable as exc:
        links, reason = [], f"index page unavailable ({exc.reason})"
    found: dict[str, _Release] = {}
    for ext in ("xlsx", "csv"):
        if candidates := [u for u in links if _is_release(u, entry.id, ext)]:
            found[ext] = _Release(max(candidates, key=release_date), stale)
        elif ext in entry.last_known:
            found[ext] = _Release(entry.last_known[ext], stale or reason)
    return found


def _is_release(url: str, prefix: str, ext: str) -> bool:
    name = url.rsplit("/", 1)[-1].lower()
    return name.startswith(f"{prefix}_") and name.endswith(f".{ext}")


def _absolute(href: str) -> str:
    return href if href.startswith("http") else urljoin(INDEX_URL, href)


def parse_workbook(payload: RawPayload, entry: DatasetEntry, platform: PlatformInfo) -> _Rows:
    """Read the classification sheet: one or more side-by-side blocks, each a "SSIC 2020" code
    column, an industry label column and period columns (quarterly blocks add a quarter row under
    the year row). Legend, unit and notes are read from the sheet itself."""
    assert entry.sheet, f"{entry.id} needs a sheet name"
    try:
        book = openpyxl.load_workbook(io.BytesIO(payload.body), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise SourceSchemaError(payload.url, f"not a readable workbook ({type(exc).__name__})") from exc
    if entry.sheet not in book.sheetnames:
        raise SourceSchemaError(
            payload.url, f"sheet {entry.sheet!r} missing (sheets: {', '.join(book.sheetnames)})"
        )
    grid = [list(r) for r in book[entry.sheet].iter_rows(values_only=True)]

    header = next((i for i, r in enumerate(grid) if any(_text(c) == BLOCK_HEADER for c in r)), None)
    if header is None:
        raise SourceSchemaError(payload.url, f"no {BLOCK_HEADER!r} header in sheet {entry.sheet!r}")
    starts = [j for j, c in enumerate(grid[header]) if _text(c) == BLOCK_HEADER]
    quarterly = entry.frequency == "quarterly"
    first_data = header + (2 if quarterly else 1)

    rows, marked_headers = _Rows(), 0
    for n, start in enumerate(starts):
        end = starts[n + 1] if n + 1 < len(starts) else len(grid[header])
        periods, marked = _block_periods(grid, header, start + 2, end, quarterly)
        marked_headers += marked
        for r in grid[first_data:]:
            label = _text(r[start + 1]) if start + 1 < len(r) else ""
            if not label:
                continue
            row = _RowId(_text(r[start]) or None, label)
            for j, period in periods:
                if j < len(r) and r[j] is not None and _text(r[j]) != "":
                    rows.points[row].append((period, _raw(r[j])))
    if not rows.points:
        raise SourceSchemaError(payload.url, f"no industry rows with period values in sheet {entry.sheet!r}")

    texts = [_text(c) for r in grid for c in r if isinstance(c, str)]
    legend = {m["marker"]: m["meaning"].strip() for t in texts if (m := _LEGEND.match(t))}
    units = sorted({m["unit"].strip().lower() for t in texts if (m := _UNIT.match(t))})
    rows.metadata = {"legend": legend, "units": units}
    rows.issues.extend(_legend_issues(legend, platform))
    rows.issues.extend(_unit_issues(units, entry))
    if marked_headers:
        rows.issues.append(
            QualityIssue(
                code="PERIOD_FOOTNOTE",
                severity=Severity.INFO,
                count=marked_headers,
                message="period headers carry a footnote mark (e.g. '1Q^') that the sheet does not explain",
            )
        )
    notes = [t for t in texts if re.match(r"^\d+\.\s", t)]
    rows.footnotes = " ".join(dict.fromkeys(notes)) or None
    return rows


def _block_periods(
    grid: list[list[Any]], header: int, first: int, end: int, quarterly: bool
) -> tuple[list[tuple[int, str]], int]:
    periods, year, marked = [], None, 0
    for j in range(first, end):
        cell = grid[header][j] if j < len(grid[header]) else None
        if isinstance(cell, int | float) and not isinstance(cell, bool):
            year = int(cell)
        if year is None:
            continue
        if not quarterly:
            if cell is not None:
                periods.append((j, str(year)))
            continue
        quarter = _text(grid[header + 1][j]) if j < len(grid[header + 1]) else ""
        if quarter:
            marked += bool(_HEADER_MARKS.search(quarter))
            periods.append((j, f"{year} {_HEADER_MARKS.sub('', quarter)}"))
    return periods, marked


def _legend_issues(legend: dict[str, str], platform: PlatformInfo) -> list[QualityIssue]:
    """The sheet's own legend must agree with the reviewed marker meanings in the manifest."""
    issues = []
    for marker, meaning in sorted(legend.items()):
        spec = platform.markers.get(marker)
        if spec is None or spec.meaning.lower() != meaning.lower():
            issues.append(
                QualityIssue(
                    code="LEGEND_CHANGED",
                    severity=Severity.WARNING,
                    message=f"sheet defines {marker!r} as {meaning!r}; the reviewed meaning is "
                    f"{spec.meaning if spec else 'none'!r}. Values with it are nulled until reviewed.",
                )
            )
    return issues


def _unit_issues(units: list[str], entry: DatasetEntry) -> list[QualityIssue]:
    """The scale must be confirmed by the sheet itself: an unstated, unrecognised or different
    unit makes every value's magnitude uncertain, so it blocks use."""
    scales = {_UNIT_SCALES.get(u.removesuffix("s")) for u in units}
    if not units or None in scales:
        stated = ", ".join(repr(u) for u in units) or "no unit"
        return [
            QualityIssue(
                code="UNIT_UNVERIFIED",
                severity=Severity.ERROR,
                message=f"sheet states {stated}; the catalogue's x{entry.scale} cannot be confirmed",
            )
        ]
    if scales != {entry.scale}:
        return [
            QualityIssue(
                code="UNIT_MISMATCH",
                severity=Severity.ERROR,
                message=f"sheet states {', '.join(units)}; catalogue expects x{entry.scale}",
            )
        ]
    return []


def parse_csv(payload: RawPayload, entry: DatasetEntry) -> _Rows:
    """The CSV: one row per (period, lowest-level industry). It states no codes, unit or legend."""
    assert entry.fields and entry.fields.period and entry.fields.value
    columns = [entry.fields.period, entry.fields.label, entry.fields.value]
    frame = require_columns(read_csv(payload), columns, payload.url)
    rows = _Rows()
    for period, label, value in frame[columns].itertuples(index=False):
        rows.points[_RowId(None, label)].append((period, value))
    return rows


def _text(cell: Any) -> str:
    return str(cell).strip() if cell is not None else ""


def _raw(cell: Any) -> str:
    """Numbers as the source wrote them (an integral float such as 14.0 as "14")."""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return _text(cell)


async def fetch_timeseries(
    entry: DatasetEntry, platform: PlatformInfo, taxonomy: Taxonomy, fetcher: Fetcher
) -> NormalizedDataset:
    releases = await _discover(entry, fetcher)
    issues: list[QualityIssue] = []
    payload: RawPayload | None = None
    rows: _Rows | None = None
    if (book := releases.get("xlsx")) is not None:
        try:
            payload = await fetcher.get(book.url, rate=POLITE, accept="*/*")
            rows = parse_workbook(payload, entry, platform)
            issues.extend(filter(None, [book.disclosure("xlsx")]))
        except (SourceUnavailable, SourceSchemaError) as exc:
            reason = exc.reason if isinstance(exc, SourceUnavailable) else exc.problem
            issues.append(
                QualityIssue(
                    code="FORMAT_FALLBACK",
                    severity=Severity.WARNING,
                    message=f"workbook unusable ({reason}); using the CSV, which lists only the "
                    "lowest-level industries (no sector totals) and states no unit (scale taken "
                    "from the catalogue)",
                )
            )
    if rows is None:
        if (table := releases.get("csv")) is None:
            raise SourceUnavailable(INDEX_URL, f"no workbook or CSV available for {entry.id}")
        payload = await fetcher.get(table.url, rate=POLITE, accept="*/*")
        rows = parse_csv(payload, entry)
        issues.extend(filter(None, [table.disclosure("csv")]))
    assert payload is not None
    issues.extend(rows.issues)

    series: list[SeriesData] = []
    for row, points in rows.points.items():
        built, label_issues = build_series(
            entry,
            platform,
            taxonomy,
            source_ref=row.ref,
            label=row.label,
            points=points,
            stated_code=row.code,
        )
        series.extend(built)
        issues.extend(label_issues)

    return assemble(
        entry,
        platform,
        taxonomy,
        [payload],
        series,
        issues,
        title=entry.title,
        source_updated_at=release_date(payload.url) or None,
        footnotes=rows.footnotes,
        source_metadata=rows.metadata,
    )
