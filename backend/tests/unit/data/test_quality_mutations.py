"""Data-quality mutation suite (plan §8.2): each test takes a REAL captured payload, breaks it in
one way sources really break, and checks the damage is detected and kept out of the evidence."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from collections.abc import Callable

import openpyxl
import pytest

from app.connectors import datagov, singstat
from app.connectors.catalog import get_catalog
from app.connectors.http import Fetcher, RawPayload, SnapshotStore, SourceSchemaError, get_snapshot_store
from app.connectors.registry import fetch
from app.data.consistency import agree, compare, tolerance
from app.data.taxonomy import get_taxonomy

CATALOG = get_catalog()
TAXONOMY = get_taxonomy()
VACANCIES_Q = "d_3f17bbf67d31d4d402caa8755eb74afc"
MOM_Q = "mrsd_3_qtly_emp_chg_by_ind"
MOM_Q_XLSX = CATALOG.get(MOM_Q).last_known["xlsx"]

Mutation = Callable[[bytes], bytes]


class MutatedSnapshots(SnapshotStore):
    """The real snapshots, with chosen payloads altered on the way out."""

    def __init__(self, mutations: dict[str, Mutation]) -> None:
        self.real, self.mutations = get_snapshot_store(), mutations

    def load(self, url: str) -> RawPayload | None:
        payload = self.real.load(url)
        if payload is None or url not in self.mutations:
            return payload
        body = self.mutations[url](payload.body)
        return payload.model_copy(update={"body": body, "sha256": hashlib.sha256(body).hexdigest()})


async def _fetch(dataset_id: str, mutations: dict[str, Mutation]):
    return await fetch(
        dataset_id, CATALOG, TAXONOMY, Fetcher(mode="snapshot", snapshots=MutatedSnapshots(mutations))
    )


def _codes(dataset) -> set[str]:
    return {i.code for i in dataset.quality}


# -- helpers that edit real payload formats ----------------------------------------------------


def _edit_csv(edit: Callable[[list[dict]], list[dict]], *, drop: str | None = None) -> Mutation:
    def mutate(body: bytes) -> bytes:
        rows = list(csv.DictReader(io.StringIO(body.decode())))
        rows = edit(rows)
        fields = [f for f in rows[0] if f != drop]
        out = io.StringIO()
        writer = csv.DictWriter(out, fieldnames=fields, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
        return out.getvalue().encode()

    return mutate


def _edit_singstat(edit: Callable[[dict], None]) -> Mutation:
    def mutate(body: bytes) -> bytes:
        page = json.loads(body)
        for row in page["Data"]["row"]:
            edit(row)
        return json.dumps(page).encode()

    return mutate


def _edit_workbook(edit: Callable[[openpyxl.worksheet.worksheet.Worksheet], None]) -> Mutation:
    def mutate(body: bytes) -> bytes:
        book = openpyxl.load_workbook(io.BytesIO(body))
        edit(book["SSIC2020"])
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    return mutate


def _cells(sheet, predicate):
    return [
        c for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and predicate(c.value.strip())
    ]


# -- the mutation classes -----------------------------------------------------------------------


async def test_dropped_column_is_schema_drift_not_a_crash():
    drop_value = _edit_csv(lambda rows: rows, drop="job_vacancy")

    def drop_from_pages(body: bytes) -> bytes:
        page = json.loads(body)
        for record in page["result"]["records"]:
            record.pop("job_vacancy", None)
        return json.dumps(page).encode()

    pages = [datagov.datastore_url(VACANCIES_Q, offset) for offset in range(0, 6000, datagov.PAGE_RECORDS)]
    mutations = {datagov.download_url(VACANCIES_Q): drop_value, **dict.fromkeys(pages, drop_from_pages)}

    with pytest.raises(SourceSchemaError, match="missing column\\(s\\) job_vacancy"):
        await _fetch(VACANCIES_Q, mutations)


async def test_dropped_column_in_the_csv_alone_falls_back_to_the_datastore():
    mutations = {datagov.download_url(VACANCIES_Q): _edit_csv(lambda rows: rows, drop="job_vacancy")}
    ds = await _fetch(VACANCIES_Q, mutations)
    assert "FORMAT_FALLBACK" in _codes(ds) and ds.find("info_comms")


async def test_renamed_label_is_disclosed_and_never_matched_to_a_sector():
    def rename(row: dict) -> None:
        if row["rowText"] == "Information And Communications":
            row["rowText"] = "Infocomm & Media"

    ds = await _fetch("M182841", {singstat.tabledata_url("M182841", 0): _edit_singstat(rename)})

    assert ds.find("info_comms") == []
    unidentified = next(i for i in ds.quality if i.code == "UNIDENTIFIED_SERIES")
    assert "Infocomm & Media" in unidentified.message


async def test_new_marker_is_nulled_and_flagged_not_parsed():
    def mark(row: dict) -> None:
        if row["rowText"] == "Information And Communications":
            row["columns"][-1]["value"] = "x"  # a symbol no legend defines

    ds = await _fetch("M182841", {singstat.tabledata_url("M182841", 0): _edit_singstat(mark)})

    (ict,) = ds.find("info_comms")
    last = ict.observations[-1]
    assert last.value is None and last.raw == "x" and last.marker == "unknown"
    assert "UNLISTED_MARKER" in _codes(ds)


async def test_duplicate_period_with_a_different_value_blocks_the_series():
    def duplicate(rows: list[dict]) -> list[dict]:
        ict = next(
            r for r in rows if r["quarter"] == "2024-Q4" and r["industry"] == "information and communications"
        )
        return [*rows, {**ict, "job_vacancy": str(int(ict["job_vacancy"]) + 500)}]

    ds = await _fetch(VACANCIES_Q, {datagov.download_url(VACANCIES_Q): _edit_csv(duplicate)})

    (ict,) = ds.find("info_comms")
    blocking = ds.blocking_issues(ict)
    assert [i.code for i in blocking] == ["CONFLICTING_DUPLICATE"] and "2024-Q4" in blocking[0].message


async def test_unit_change_in_the_workbook_is_an_error():
    def hundreds(sheet) -> None:
        for cell in _cells(sheet, lambda v: v == "In Thousands"):
            cell.value = "In Hundreds"

    ds = await _fetch(MOM_Q, {MOM_Q_XLSX: _edit_workbook(hundreds)})

    mismatch = next(i for i in ds.quality if i.code == "UNIT_MISMATCH")
    assert mismatch.severity == "error" and "x1000" in mismatch.message
    (ict,) = ds.find("info_comms")
    assert ds.blocking_issues(ict), "evidence from a dataset with the wrong scale is unusable"


async def test_child_relabelled_as_its_parent_is_an_identity_conflict():
    def relabel(sheet) -> None:
        for cell in _cells(sheet, lambda v: v == "IT & Other Information Services"):
            cell.value = "INFORMATION AND COMMUNICATIONS"  # the J62-63 row now claims the parent's name

    ds = await _fetch(MOM_Q, {MOM_Q_XLSX: _edit_workbook(relabel)})

    conflict = next(i for i in ds.quality if i.code == "IDENTITY_CONFLICT")
    assert "J62-63" in conflict.message and conflict.severity == "error"
    (ict,) = ds.find("info_comms")
    assert ict.source_ref.startswith("J58-63"), "the genuine parent row is the one identified"
    assert ds.find("it_info_services") == []


async def test_singstat_child_relabelled_as_parent_makes_both_ambiguous():
    def relabel(row: dict) -> None:
        if row["seriesNo"] == "1.2.4.2":  # IT & Other Information Services
            row["rowText"] = "Information And Communications"

    ds = await _fetch("M184071", {singstat.tabledata_url("M184071", 0): _edit_singstat(relabel)})

    assert ds.find("info_comms") == [] and "AMBIGUOUS_SECTOR" in _codes(ds)


async def test_changed_legend_is_flagged():
    def redefine(sheet) -> None:
        for cell in _cells(sheet, lambda v: v.startswith("'-'")):
            cell.value = "'-' :  Not available"

    ds = await _fetch(MOM_Q, {MOM_Q_XLSX: _edit_workbook(redefine)})

    assert "LEGEND_CHANGED" in _codes(ds)


async def test_missing_sheet_falls_back_to_the_csv_and_says_what_is_lost():
    def rename_sheet(body: bytes) -> bytes:
        book = openpyxl.load_workbook(io.BytesIO(body))
        book["SSIC2020"].title = "SSIC2025"
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    ds = await _fetch(MOM_Q, {MOM_Q_XLSX: rename_sheet})

    fallback = next(i for i in ds.quality if i.code == "FORMAT_FALLBACK")
    assert "SSIC2020" in fallback.message and "no sector totals" in fallback.message
    assert ds.find("info_comms") == [], "the CSV has no Information & Communications total"
    (it,) = ds.find("it_info_services")
    assert it.value_map()["2024-Q4"].value == -1.0


async def test_mirror_mismatch_is_reported_period_by_period():
    def alter(body: bytes) -> bytes:
        return _edit_csv(
            lambda rows: [
                {**r, "job_vacancy": "9999"}
                if r["quarter"] == "2023-Q2" and r["industry"] == "information and communications"
                else r
                for r in rows
            ]
        )(body)

    singstat_ds = await _fetch("M184071", {})
    mirror_ds = await _fetch(VACANCIES_Q, {datagov.download_url(VACANCIES_Q): alter})
    a, b = singstat_ds.find("info_comms")[0], mirror_ds.find("info_comms")[0]
    window = [f"{y}-Q{q}" for y in range(2020, 2025) for q in range(1, 5)]

    compared, mismatches = compare(a, b, window, tolerance(CATALOG.get("M184071"), CATALOG.get(VACANCIES_Q)))

    assert len(compared) == 20
    assert mismatches == [{"period": "2023-Q2", a.key: a.value_map()["2023-Q2"].value, b.key: 9999.0}]


async def test_unmutated_mirrors_agree_including_across_published_precision():
    """Control: MOM publishes thousands to 0.1, SingStat persons; they agree once rounding is allowed."""
    singstat_ds, mom_ds = await _fetch("M183891", {}), await _fetch(MOM_Q, {})
    a, b = singstat_ds.find("info_comms")[0], mom_ds.find("info_comms")[0]
    periods = sorted(set(a.value_map()) & set(b.value_map()))
    allowed = tolerance(CATALOG.get("M183891"), CATALOG.get(MOM_Q))

    compared, mismatches = compare(a, b, periods, allowed)

    assert allowed == 50 and len(compared) >= 36 and mismatches == []
    # Rounding is allowed for, real differences are not: 0.1 thousand apart is a mismatch.
    assert not agree(-1100.0, -1.2 * 1000, allowed)


@pytest.mark.parametrize("unit", ["In Millions", None])
async def test_unrecognised_or_missing_workbook_unit_blocks_use(unit):
    def change(sheet) -> None:
        for cell in _cells(sheet, lambda v: v == "In Thousands"):
            cell.value = unit

    ds = await _fetch(MOM_Q, {MOM_Q_XLSX: _edit_workbook(change)})

    (ict,) = ds.find("info_comms")
    assert [i.code for i in ds.blocking_issues(ict)] == ["UNIT_UNVERIFIED"]


async def test_a_changed_unit_is_different_evidence_with_its_own_identity():
    def hundreds(sheet) -> None:
        for cell in _cells(sheet, lambda v: v == "In Thousands"):
            cell.value = "In Hundreds"

    clean, changed = await _fetch(MOM_Q, {}), await _fetch(MOM_Q, {MOM_Q_XLSX: _edit_workbook(hundreds)})

    assert [s.model_dump() for s in clean.series] == [s.model_dump() for s in changed.series]
    assert clean.artifact_id != changed.artifact_id, "same numbers, different stated unit"
